"""Reading a QuickBooks Online export (`IMP-01`, `IMP-02`).

A QuickBooks Online export is a zip of `.xlsx` reports. This reads four of them and produces
`SourceBooks`, so nothing below it learns what QuickBooks is and a second source is a reader
rather than a second pipeline.

**The reader transforms inputs only.** ADR-0010 rejected a Beancount translator because it
would be "untested code sitting between the system and its own correctness evidence… a defect in
it produces false agreement as readily as false divergence". That objection does not reach here:
the expected answer arrives through a *different report in the same export*, so a reader bug
corrupts the journal and the reconciliation then fails loudly rather than agreeing falsely.

**Account types come from the source's own statements**, not from us. An account under `ASSETS`
on the balance sheet is an asset; one under `Income` on the profit and loss is income. Guessing
a type from an account's name would put our judgement into data whose whole value is that it is
not ours.

**Accounts are identified by name.** QuickBooks has no account codes, so the name — the full
colon-separated path — is what the two systems agree on, and it is what the `code` carries here.

**Basis matters, and the export's reports are not all usable.** Every report in a QuickBooks
export carries its accounting method in its footer. `Journal.xlsx` is the accrual record; a
report run on a cash basis does not sum to it and must not be used as its oracle.

**Figures in the statement reports are formulas, and a formula is not a zero.** QuickBooks
writes `Balance_sheet.xlsx`, `Profit_and_loss.xlsx` and `Trial_balance.xlsx` with every amount
as `=2492948.04` and no cached result, while `Journal.xlsx` and `General_ledger.xlsx` carry
literal numbers. A reader asking openpyxl for cached values gets `0.0` for every figure in the
first three and correct values from the last two — silently, and in the direction that looks
like a balanced book. So the workbook is read unevaluated, a bare numeric formula is the number
it states, and anything else is refused rather than guessed.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from collections.abc import Sequence
from datetime import date, datetime
from decimal import Decimal

import openpyxl

from cfokit.imports.source import (
    SourceAccount,
    SourceBooks,
    SourceEntry,
    SourceLine,
    StatedBalance,
    StatedStatement,
)
from cfokit.ledger.errors import LedgerError

__all__ = [
    "MAX_ARCHIVE_BYTES",
    "MAX_MEMBERS",
    "MAX_UNCOMPRESSED_BYTES",
    "SYSTEM",
    "ImportTooLarge",
    "SourceBooks",
    "UnreadableFigure",
    "read",
    "read_basis",
]

# What `IMP-04` requires an imported record to name.
SYSTEM = "QuickBooks Online"

# A formula whose whole body is the number it states, which is how QuickBooks writes a figure.
_LITERAL = re.compile(r"-?\d+(\.\d+)?")


class UnreadableFigure(LedgerError):
    """A cell whose value cannot be established, in a place a figure belongs.

    Refused rather than defaulted. A missing amount that arrives as zero is indistinguishable
    from a real zero, and in a set of books the difference is everything.
    """

    code = "unreadable_figure"
    status = 422


# QuickBooks Online exports one currency per company file, and states it nowhere in these
# reports. Declared here rather than guessed from a cell: `IMP-07` refuses a foreign amount, and
# that refusal is worth nothing if the commodity was inferred.
_COMMODITY = "USD"

Row = tuple[object, ...]

# Column positions in the Journal report, from its header row.
_DATE, _ACCOUNT, _DEBIT, _CREDIT, _MEMO = 1, 6, 7, 8, 5
_HEADER_ROWS = 5

_BALANCE_SHEET_SECTIONS = {
    "ASSETS": "asset",
    "LIABILITIES AND EQUITY": None,  # split further by its own subsections
    "Liabilities": "liability",
    "Equity": "equity",
}
# A profit and loss has four sections, not two. "Other Income" and "Other Expenses" carry
# things outside the trading result — interest earned, a gain on disposal — and an account
# under one of them is still income or expense. Reading only the first two leaves those
# accounts typed `unknown`, which lands them in the chart as assets and then compares them
# against the statement with the wrong sign.
_INCOME_STATEMENT_SECTIONS = {
    "Income": "income",
    "Other Income": "income",
    "Expenses": "expense",
    "Other Expenses": "expense",
}


def read_basis(rows: Sequence[Row]) -> str:
    """The accounting method a report was run on, from its footer.

    QuickBooks prints it on the last labelled row — "… - Cash Basis" or "… - Accrual Basis".
    A report whose basis cannot be read is `unknown`, which callers must treat as unusable
    rather than assume.
    """
    for row in reversed(rows):
        label = row[0]
        if label is None:
            continue
        text = str(label)
        if "Cash Basis" in text:
            return "cash"
        if "Accrual Basis" in text:
            return "accrual"
        break
    return "unknown"


class ImportTooLarge(LedgerError):
    """An archive this surface will not read, refused before it is decompressed.

    **A ceiling belongs here, not at an adapter.** It was at the MCP tool, so the CLI had none
    and any endpoint added later would have inherited none either. Every path onto the books
    goes through `read`, which is the only place holding the bytes.

    Three bounds rather than one, because one only catches the honest mistake:

    - `MAX_ARCHIVE_BYTES` on the file, which fails fastest and costs nothing to check.
    - `MAX_UNCOMPRESSED_BYTES` on what it expands to, which is the one that matters. A zip's
      compressed size bounds its expanded size only if you assume the ratio is sane, and an
      archive is untrusted input: 1,029x is four lines of Python, so a file inside a 100 MB
      ceiling expands past 100 GB and takes the process — and every other entity's tools in it
      — with it (`NFR-04`).
    - `MAX_MEMBERS`, because thousands of tiny members cost time and file handles without
      tripping either size bound.
    """

    code = "import_too_large"
    status = 413


# Measured against a real export rather than guessed: eight members, 979,403 bytes expanded
# from 904,749, an overall ratio of 1.08. That ratio is not luck — the members are `.xlsx`
# files, which are themselves zips, so a genuine export is very nearly incompressible. These
# leave room for an export a hundred times larger than the one seen here while refusing
# anything whose ratio says it is not an accounting export at all.
MAX_ARCHIVE_BYTES = 100 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 300 * 1024 * 1024
MAX_MEMBERS = 64


def _bounded(export: zipfile.ZipFile) -> None:
    """Refuse an archive that would expand past what this surface reads.

    Checked against the central directory, which is metadata rather than decompression, so a
    bomb is refused without a byte of it being expanded. The directory is the attacker's to
    write, so it is not trusted on its own — `_member` bounds the actual read as well. This
    check exists so the ordinary case fails immediately with a message naming the limit.
    """
    members = export.infolist()
    if len(members) > MAX_MEMBERS:
        raise ImportTooLarge(
            f"the archive holds {len(members)} members; this surface reads at most"
            f" {MAX_MEMBERS}"
        )
    declared = sum(member.file_size for member in members)
    if declared > MAX_UNCOMPRESSED_BYTES:
        raise ImportTooLarge(
            f"the archive declares {declared} bytes expanded; this surface reads at most"
            f" {MAX_UNCOMPRESSED_BYTES}"
        )


def _member(export: zipfile.ZipFile, name: str) -> bytes:
    """One member's bytes, refusing to read past the ceiling.

    Streamed with an explicit stop rather than `ZipFile.read`, because the size `_bounded`
    checked came from the central directory and an attacker writes that. A member declaring a
    kilobyte and delivering a terabyte passes the metadata check and is stopped here.
    """
    with export.open(name) as stream:
        payload = stream.read(MAX_UNCOMPRESSED_BYTES + 1)
    if len(payload) > MAX_UNCOMPRESSED_BYTES:
        raise ImportTooLarge(
            f"{name} expands past {MAX_UNCOMPRESSED_BYTES} bytes; the archive's own"
            " directory understated it"
        )
    return payload


def read(archive: bytes) -> SourceBooks:
    """One QuickBooks export, as `SourceBooks`.

    The general ledger's own per-account totals are the oracle, not the trial balance. Every
    *report* in an export carries its accounting method in its footer; `Journal.xlsx` carries
    none, because it is the raw record rather than a view of it.

    So the two bases come from different files and mean different things. `basis` is the
    journal's — `unknown` for QuickBooks, which is honest: it prints none. `balances_basis` is
    the general ledger's, and it is routinely `cash` on the same export whose journal is
    accrual. Reconciling one against the other then diverges on exactly the obligation
    accounts, which ADR-0037 predicts and `IMP-08` reports rather than tolerates.
    """
    if len(archive) > MAX_ARCHIVE_BYTES:
        raise ImportTooLarge(
            f"the archive is {len(archive)} bytes; this surface reads at most"
            f" {MAX_ARCHIVE_BYTES}"
        )
    with zipfile.ZipFile(io.BytesIO(archive)) as export:
        _bounded(export)
        types = _account_types(export)
        entries = _entries(export)
        basis = read_basis(_rows(export, "Journal.xlsx"))
        balances, rollups, balances_basis = _ledger_totals(export, entries)
        commodity = _COMMODITY
        posted_to = {line.account_code for entry in entries for line in entry.lines}
        printed = tuple(
            statement
            for member, report in (
                ("Profit_and_loss.xlsx", "profit_and_loss"),
                ("Balance_sheet.xlsx", "balance_sheet"),
            )
            if (statement := _statement(export, member, report, posted_to)) is not None
        )

    used = sorted({line.account_code for entry in entries for line in entry.lines})
    return SourceBooks(
        system=SYSTEM,
        # Over the archive bytes, so re-reading the same export yields the same identity and a
        # different export never collides with it. Not a security boundary — nobody is
        # adversarially colliding their own accounting export — so sha256 here is simply the
        # obvious digest rather than a considered choice against a weaker one.
        fingerprint=hashlib.sha256(archive).hexdigest(),
        basis=basis,
        balances_basis=balances_basis,
        commodity=commodity,
        accounts=tuple(
            SourceAccount(
                code=name,
                name=name,
                account_type=_typed(name, types),
                parent=_parent(name, set(used)),
            )
            for name in used
        ),
        entries=entries,
        balances=balances,
        rollups=rollups,
        statements=printed,
    )


def _statement(
    export: zipfile.ZipFile, member: str, report: str, posted_to: set[str]
) -> StatedStatement | None:
    """One printed statement, by account.

    **Account paths come from the indentation.** A statement indents a sub-account under its
    parent, so `Business Development` beneath `Advertising & Marketing` is the account the
    journal calls `Advertising & Marketing:Business Development`. Reading the leaf name alone
    would be ambiguous — `Meals` appears under two different parents in a real chart.

    But a statement's hierarchy is not only accounts: a balance sheet groups them under
    `ASSETS`, `Current Assets`, `Bank Accounts`, none of which anything posts to. So the path
    the indentation suggests is *checked* against the accounts the journal uses, longest first,
    and a row that matches none is reported as unmatched rather than guessed at.

    Sign is left exactly as the source prints it. A statement shows income as positive; a
    posting signs income negative. Translating here would bury the convention in a reader —
    the comparison is where two conventions meet, so that is where it belongs.
    """
    if member not in export.namelist():
        return None
    rows = _rows(export, member)
    lines: list[StatedBalance] = []
    unmatched: list[str] = []
    stack: list[tuple[int, str]] = []

    for row in rows[_HEADER_ROWS:]:
        label = row[0]
        if label is None:
            continue
        text = str(label)
        name = text.strip()
        if not name or name.upper().startswith("TOTAL") or name.startswith("Gross "):
            continue
        if "Basis" in name and name.endswith("Basis"):
            continue

        depth = len(text) - len(text.lstrip())
        while stack and stack[-1][0] >= depth:
            stack.pop()
        stack.append((depth, name))

        if row[1] is None:
            continue  # a heading, or a parent that carries no figure of its own

        code = _resolve([entry for _, entry in stack], posted_to)
        if code is None:
            unmatched.append(name)
            continue
        lines.append(StatedBalance(account_code=code, balance=_amount(row[1])))

    return StatedStatement(
        report=report,
        basis=read_basis(rows),
        lines=tuple(lines),
        unmatched=tuple(unmatched),
    )


def _resolve(path: list[str], posted_to: set[str]) -> str | None:
    """The account a statement row names, from the path its indentation implies.

    Longest suffix first, because the deeper path is the more specific claim: a row nested
    under a real parent is that parent's sub-account, and one nested only under grouping
    headings is a top-level account whose leaf name is the whole of it.
    """
    for start in range(len(path)):
        candidate = ":".join(path[start:])
        if candidate in posted_to:
            return candidate
    return None


def _typed(name: str, types: dict[str, str]) -> str:
    """An account's type, from the statement it appears on or from its parent's.

    QuickBooks prints a sub-account under its bare leaf name on the balance sheet and the profit
    and loss, so the full colon-separated path the journal uses never matches one of those rows.
    Walking up the path is not a guess about the account: QuickBooks requires a sub-account to
    share its parent's type, and the path is the source stating the hierarchy.

    Unknown means the account appears in the journal and beneath no statement section at all —
    recorded as such rather than guessed, because a wrong type is a silent misclassification and
    a missing one is a question somebody can answer.
    """
    parts = name.split(":")
    for cut in range(len(parts), 0, -1):
        stated = types.get(":".join(parts[:cut]))
        if stated is not None:
            return stated
    # Last resort: the bare leaf. A statement prints "Accounting Services" indented under its
    # parent's header, and the header itself carries no amount when the parent takes no
    # postings of its own — so neither the full path nor any prefix of it matches a row, while
    # the leaf matches exactly. Tried after every prefix, because a leaf name can repeat under
    # two parents and the more specific answer should win where there is one.
    return types.get(parts[-1], "unknown")


def _parent(name: str, accounts: set[str]) -> str:
    """The account this one sits under, where the export has one.

    A parent that takes no postings of its own is not in the journal and so is not an account
    here; the child is then reported as top-level rather than pointing at something absent.
    """
    parts = name.split(":")
    for cut in range(len(parts) - 1, 0, -1):
        candidate = ":".join(parts[:cut])
        if candidate in accounts:
            return candidate
    return ""


def _account_types(export: zipfile.ZipFile) -> dict[str, str]:
    """Map each account to a type, using the source's own statement sections."""
    types: dict[str, str] = {}
    for member, sections in (
        ("Balance_sheet.xlsx", _BALANCE_SHEET_SECTIONS),
        ("Profit_and_loss.xlsx", _INCOME_STATEMENT_SECTIONS),
    ):
        if member not in export.namelist():
            continue
        rows = _rows(export, member)
        current: str | None = None
        for row in rows[_HEADER_ROWS:]:
            label = row[0]
            if label is None:
                continue
            text = str(label).strip()
            if text in sections:
                current = sections[text]
                continue
            # "Total X" and "TOTAL X" are subtotals, not accounts.
            if text.upper().startswith("TOTAL") or text.startswith("Gross "):
                continue
            if current is not None and any(cell is not None for cell in row[1:]):
                types.setdefault(text, current)
    return types


def _entries(export: zipfile.ZipFile) -> tuple[SourceEntry, ...]:
    """Every transaction in the journal, with its lines.

    A transaction begins on the row carrying a date; its remaining lines carry none. Each group
    ends with a totals row that has amounts and no account — skipped, or every transaction would
    count twice.

    The reference is positional because QuickBooks' journal prints no stable identifier for a
    transaction. Positional is honest about what it is: it identifies the row in *this* export,
    which is what a person tracing a posting back to the file needs.
    """
    rows = _rows(export, "Journal.xlsx")
    entries: list[SourceEntry] = []
    lines: list[SourceLine] = []
    reference = 0
    when: date | None = None
    description = ""

    def flush() -> None:
        if when is not None:
            entries.append(
                SourceEntry(
                    reference=str(reference),
                    transaction_date=when,
                    description=description,
                    lines=tuple(lines),
                )
            )

    for row in rows[_HEADER_ROWS:]:
        if row[_DATE] is not None:
            flush()
            lines = []
            reference += 1
            when = _as_date(row[_DATE])
            description = str(row[_MEMO] or "")
        account = row[_ACCOUNT]
        if account is None or when is None:
            continue
        lines.append(
            SourceLine(
                account_code=str(account),
                amount=_amount(row[_DEBIT]) - _amount(row[_CREDIT]),
                commodity=_COMMODITY,
            )
        )
    flush()
    return tuple(entries)


def _ledger_totals(
    export: zipfile.ZipFile, entries: Sequence[SourceEntry]
) -> tuple[tuple[StatedBalance, ...], tuple[StatedBalance, ...], str]:
    """The general ledger's own stated totals, split into balances and rollups, and its basis.

    QuickBooks closes each account's section with a "Total for <account>" row. Those are the
    source's own balances — the figure `IMP-08` reconciles against — and they cover only
    accounts with activity, which is why an account may be in the journal and absent here.

    Some of those rows are **subtotals over a parent and its children**, printed as
    "Total for X with sub-accounts" where the parent also takes postings of its own, and as a
    bare "Total for X" where it does not. Both are recognised, and the bare case is decided
    against the export's own data rather than its wording: a bare total is a rollup only when
    the journal posts beneath that parent and never to the parent itself. A parent that takes
    direct postings therefore keeps its own balance and contributes its "with sub-accounts" row
    to the rollups, which is what those two rows mean.

    They are held apart because comparing a subtotal against a chart account would manufacture
    a divergence out of a category error.
    """
    if "General_ledger.xlsx" not in export.namelist():
        return (), (), "unknown"
    rows = _rows(export, "General_ledger.xlsx")
    posted_to = {line.account_code for entry in entries for line in entry.lines}
    balances: list[StatedBalance] = []
    rollups: list[StatedBalance] = []
    for row in rows[_HEADER_ROWS:]:
        label = row[0]
        if label is None or not str(label).startswith("Total for "):
            continue
        account = str(label)[len("Total for ") :].strip()
        parent = account.removesuffix(" with sub-accounts")
        stated = StatedBalance(
            account_code=parent, balance=_amount(row[_DEBIT]) - _amount(row[_CREDIT])
        )
        has_children = any(code.startswith(parent + ":") for code in posted_to)
        if account != parent or (has_children and parent not in posted_to):
            rollups.append(stated)
        else:
            balances.append(stated)
    return tuple(balances), tuple(rollups), read_basis(rows)


def _rows(export: zipfile.ZipFile, member: str) -> list[Row]:
    """One report's cells.

    Rows rather than a worksheet, so nothing downstream needs openpyxl's types — and the two
    shapes `load_workbook` returns for read-only and normal mode stop mattering.
    """
    # Unevaluated. `data_only=True` asks for a cached result, and QuickBooks caches none — so
    # every figure in the statement reports would come back as 0.0 rather than as missing.
    workbook = openpyxl.load_workbook(
        io.BytesIO(_member(export, member)), read_only=True, data_only=False
    )
    try:
        return list(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()


def _amount(value: object) -> Decimal:
    """A cell as an exact decimal.

    `Decimal(str(value))` rather than `Decimal(value)`: a float's binary expansion is not the
    figure the workbook shows, and `LED-04` allows no representation error.

    **A formula that states a number is that number; any other formula is refused.** The
    statement reports write every figure as `=2492948.04`, so reading them at all means reading
    those. A formula that references other cells is a subtotal, and computing it would mean
    implementing a spreadsheet — where returning zero instead would be worse than either,
    because a zero in a financial figure reads as a fact.
    """
    if value is None:
        return Decimal(0)
    if isinstance(value, str) and value.startswith("="):
        stated = value[1:].strip()
        if not _LITERAL.fullmatch(stated):
            raise UnreadableFigure(f"{value[:40]!r} is a computed cell, not a stated figure")
        return Decimal(stated)
    return Decimal(str(value))


def _as_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value), "%m/%d/%Y").date()  # noqa: DTZ007
